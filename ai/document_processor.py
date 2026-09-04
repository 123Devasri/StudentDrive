import os
import pymupdf  # PyMuPDF
import docx
from pptx import Presentation

def extract_text_from_file(file_path: str) -> list:
    """
    Extracts text from PDF, PPTX, DOCX, TXT, or MD files.
    Returns a list of dicts: [{'page_number': int|None, 'slide_number': int|None, 'page': int, 'text': str}]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    pages = []

    if ext == '.pdf':
        doc = pymupdf.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()
            if text:
                pages.append({
                    'page': page_num + 1,
                    'page_number': page_num + 1,
                    'slide_number': None,
                    'text': text
                })
        doc.close()

    elif ext in ('.ppt', '.pptx'):
        prs = Presentation(file_path)
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        slide_texts.append(t)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            slide_texts.append(" | ".join(row_text))
                elif hasattr(shape, "text") and shape.text:
                    t = shape.text.strip()
                    if t:
                        slide_texts.append(t)
            if slide_texts:
                pages.append({
                    'page': slide_idx,
                    'page_number': None,
                    'slide_number': slide_idx,
                    'text': "\n".join(slide_texts)
                })

    elif ext in ('.doc', '.docx'):
        doc = docx.Document(file_path)
        doc_texts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                doc_texts.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    doc_texts.append(" | ".join(row_texts))
        if doc_texts:
            pages.append({
                'page': 1,
                'page_number': 1,
                'slide_number': None,
                'text': "\n".join(doc_texts)
            })

    elif ext in ('.txt', '.md'):
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read().strip()
            if content:
                pages.append({
                    'page': 1,
                    'page_number': 1,
                    'slide_number': None,
                    'text': content
                })
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return pages
